"""
Document skill: write Word documents from a natural-language brief.

The LLM is used to draft the body text, then python-docx writes it as a
proper .docx file. The file is saved under ~/Documents/CHAPPIE and (by
default) opened immediately in the user's editor.
"""

import os
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional

from core.skills import Skill
from utils.logger import setup_logger

logger = setup_logger(__name__)


SAFE_NAME_RX = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


class DocumentSkill(Skill):
    name = "document"
    description = (
        "Write a Word document — letter, memo, essay, draft, report. "
        "Pass a topic and optional recipient. Returns the saved .docx path."
    )
    triggers = [
        # "write a letter to John about the meeting tomorrow"
        r"^write\s+(?:a\s+|an\s+)?(?P<doc_type>letter|memo|essay|note|document|report|article|draft|cover\s*letter|resignation\s*letter|thank\s*you\s*note)"
        r"(?:\s+to\s+(?P<recipient>[^,]+?))?"
        r"(?:\s+about\s+(?P<topic>.+?))?\s*$",
        # "create a word document called Resignation"
        r"^(?:create|make|new|start)\s+(?:a\s+|an\s+)?(?:word\s+)?(?:doc|document|file)"
        r"\s+(?:called\s+|named\s+|titled\s+)?(?P<name>.+?)\s*$",
        # "draft a quick X about Y"
        r"^draft\s+(?:a\s+|an\s+)?(?P<doc_type>quick\s+)?(?P<topic>.+?)\s*$",
    ]

    def __init__(self, kb, config: Dict, llm=None):
        super().__init__(kb, config, llm)
        skills_cfg = config.get("skills", {})
        self.save_dir = os.path.expanduser(
            skills_cfg.get("document_dir", "~/Documents/CHAPPIE")
        )
        self.auto_open = skills_cfg.get("document_auto_open", True)
        Path(self.save_dir).mkdir(parents=True, exist_ok=True)

    async def execute(self, params: Dict, context: Optional[Dict] = None) -> Dict:
        try:
            from docx import Document  # lazy import
        except ImportError:
            return {
                "success": False,
                "skill": self.name,
                "message": (
                    "I need python-docx to write Word documents. "
                    "Run: pip install python-docx"
                ),
            }

        topic = (params.get("topic") or "").strip()
        recipient = (params.get("recipient") or "").strip()
        name = (params.get("name") or "").strip()
        doc_type = (params.get("doc_type") or "letter").strip().lower()
        body_override = (params.get("body") or "").strip()

        # Sensible filename
        if not name:
            if recipient:
                name = f"{doc_type.title()} to {recipient}"
            elif topic:
                name = topic[:60]
            else:
                name = f"Document {datetime.now().strftime('%Y-%m-%d_%H%M')}"

        safe_name = SAFE_NAME_RX.sub("_", name).strip().rstrip(".")
        path = os.path.join(self.save_dir, f"{safe_name}.docx")

        # Build the body via the LLM, unless the caller supplied one.
        body = body_override
        if not body and self.llm and (topic or recipient or name):
            body = await self._draft_body(doc_type, recipient, topic, name)

        if not body:
            body = f"[{doc_type.capitalize()}]\n\n"  # empty stub

        # Write it
        doc = Document()
        if recipient:
            doc.add_heading(f"To: {recipient}", level=2)
        if topic:
            doc.add_heading(topic, level=1)
        elif name and not recipient:
            doc.add_heading(name, level=1)

        for paragraph in body.split("\n\n"):
            doc.add_paragraph(paragraph.strip())

        doc.save(path)
        self.kb.track_resource(path, "document")
        logger.info(f"Wrote document: {path}")

        opened = ""
        if self.auto_open:
            opened = self._open_file(path)

        return {
            "success": True,
            "skill": self.name,
            "action": "document_create",
            "path": path,
            "name": safe_name,
            "message": f"Saved '{safe_name}.docx' to {self.save_dir}.{opened}",
        }

    async def _draft_body(
        self, doc_type: str, recipient: str, topic: str, name: str
    ) -> str:
        prompt_parts = [f"Draft a {doc_type}"]
        if recipient:
            prompt_parts.append(f"to {recipient}")
        if topic:
            prompt_parts.append(f"about {topic}")
        if not recipient and not topic and name:
            prompt_parts.append(f"on the topic of: {name}")
        prompt = " ".join(prompt_parts) + "."

        try:
            text = await self.llm.generate(
                messages=[{"role": "user", "content": prompt}],
                system=(
                    "You are a skilled writer. Produce only the document body — "
                    "no preamble, no 'Sure, here is', no 'As an AI'. "
                    "Use natural paragraph breaks. Sign off with [Your name]."
                ),
                temperature=0.6,
            )
            return text or ""
        except Exception as exc:
            logger.warning(f"LLM draft failed: {exc}")
            return ""

    def _open_file(self, path: str) -> str:
        try:
            if sys.platform == "win32":
                os.startfile(path)  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.Popen(["open", path])
            else:
                subprocess.Popen(["xdg-open", path])
            return " I've opened it for you."
        except Exception as exc:
            logger.warning(f"Couldn't auto-open document: {exc}")
            return ""
